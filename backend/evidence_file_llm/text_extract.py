"""从证据附件字节中提取可供大模型分析的纯文本（与 OCR / PDF / Office 等解耦）。"""

from __future__ import annotations

import io
import re
from pathlib import Path

from evidence_common import allowed_evidence_attachment, allowed_evidence_image

_MAX_PDF_PAGES = 30
_MAX_CHARS = 28000


def _truncate(s: str) -> str:
    s = (s or "").strip()
    if len(s) <= _MAX_CHARS:
        return s
    return s[:_MAX_CHARS] + "\n…（内容过长已截断）"


def _extract_pdf(raw: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(raw))
    parts: list[str] = []
    n = min(len(reader.pages), _MAX_PDF_PAGES)
    for i in range(n):
        try:
            t = reader.pages[i].extract_text() or ""
        except Exception:
            t = ""
        if t.strip():
            parts.append(t)
    return "\n".join(parts)


def _extract_docx(raw: bytes) -> str:
    import docx

    doc = docx.Document(io.BytesIO(raw))
    return "\n".join(p.text for p in doc.paragraphs if p.text and p.text.strip())


def _extract_xlsx(raw: bytes) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    ws = wb.active
    rows: list[str] = []
    for row in ws.iter_rows(max_row=500, max_col=50, values_only=True):
        cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
        if cells:
            rows.append("\t".join(cells))
    wb.close()
    return "\n".join(rows)


def _extract_txt(raw: bytes) -> str:
    for enc in ("utf-8", "utf-8-sig", "gbk", "gb18030"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def extract_text_from_bytes(
    filename: str,
    content_type: str,
    raw: bytes,
    *,
    ocr_image_fn,
) -> str:
    """
    ocr_image_fn: (b64: str, lang: str) -> tuple[str, dict] 与 youdao recognize_base64 一致。
    """
    if not raw:
        return ""

    fn = filename or ""
    suf = Path(fn).suffix.lower()
    ct = (content_type or "").split(";")[0].strip().lower()

    if allowed_evidence_image(fn, ct):
        import base64

        b64 = base64.b64encode(raw).decode("ascii")
        text, _ = ocr_image_fn(b64, "zh-CHS")
        return _truncate(text or "")

    if suf == ".pdf" or ct == "application/pdf":
        try:
            return _truncate(_extract_pdf(raw))
        except Exception as exc:
            raise ValueError(f"无法读取 PDF 文本：{exc}") from exc

    if suf == ".docx" or ct == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        try:
            return _truncate(_extract_docx(raw))
        except Exception as exc:
            raise ValueError(f"无法读取 Word 文档：{exc}") from exc

    if suf == ".xlsx" or ct == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        try:
            return _truncate(_extract_xlsx(raw))
        except Exception as exc:
            raise ValueError(f"无法读取 Excel：{exc}") from exc

    if suf == ".xls" or ct == "application/vnd.ms-excel":
        raise ValueError("暂不支持 .xls，请另存为 .xlsx 或 PDF 后上传")

    if suf == ".doc" or ct == "application/msword":
        raise ValueError("暂不支持 .doc，请另存为 .docx 或 PDF 后上传")

    if suf == ".txt" or ct == "text/plain":
        return _truncate(_extract_txt(raw))

    if allowed_evidence_attachment(fn, ct):
        # 兜底：尝试按 UTF-8 文本读
        try:
            t = _extract_txt(raw)
            if re.search(r"[\u4e00-\u9fff]", t) or len(t) > 200:
                return _truncate(t)
        except Exception:
            pass
        raise ValueError("暂不支持从该文件自动提取文本，请上传图片、PDF、Word、Excel 或 TXT")

    raise ValueError("不支持的文件类型")
